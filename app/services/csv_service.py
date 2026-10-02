import csv
import io
import re
import urllib.parse
from app.extensions import db
from app.models import Player, PlayerRole, PlayerCategory, PlayerStatus
from app.models.setting import SystemSettings
from app.models.audit import AuditLog

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# The 7 canonical fields strictly required by the auction system
TEMPLATE_COLUMNS = [
    'Roll Number',
    'Name',
    'Photo',
    'Year',
    'Role',
    'Branch',
    'Category'
]

HEADER_ALIASES = {
    'rule_number': [
        'rule number', 'rule no', 'rule_no', 'ruleno', 'rule_number',
        'roll number', 'roll no', 'roll_no', 'rollno', 'roll_number', 'player id', 'roll'
    ],
    'name': ['name', 'player name', 'player_name', 'full name', 'fullname'],
    'photo': ['photo', 'photo url', 'photo_url', 'image', 'image_url', 'picture', 'pic'],
    'year': ['year', 'current year', 'batch', 'study year', 'class year', 'yr'],
    'role': ['role', 'player role', 'player_role', 'playing role'],
    'branch': ['branch', 'department', 'dept', 'stream'],
    'category': ['category', 'player category', 'tier', 'grade']
}

def normalize_photo_url(url):
    """
    Transforms any photo URL, converting Google Drive sharing/view/open links
    into direct image stream URLs (via Google CDN lh3).
    """
    if not url or not isinstance(url, str):
        return url
    u = url.strip()
    if not u:
        return u
    if any(k in u for k in ['drive.google.com', 'docs.google.com', 'drive.usercontent.google.com', 'lh3.googleusercontent.com']):
        m = re.search(r'/d/([a-zA-Z0-9_-]+)', u)
        if m:
            return f"https://lh3.googleusercontent.com/d/{m.group(1)}"
        m = re.search(r'[?&]id=([a-zA-Z0-9_-]+)', u)
        if m:
            return f"https://lh3.googleusercontent.com/d/{m.group(1)}"
    return u

def is_valid_image_url(url):
    """Validate whether a string is a well-formed http/https image URL or Drive link."""
    if not url or not isinstance(url, str):
        return False
    u = url.strip()
    if not (u.startswith('http://') or u.startswith('https://')):
        return False
    try:
        parsed = urllib.parse.urlparse(u)
        return bool(parsed.scheme and parsed.netloc)
    except Exception:
        return False

def normalize_year(val):
    """Normalize year input to a clean string ('1', '2', '3', '4', etc.)."""
    if val is None:
        return ''
    s = str(val).strip()
    if not s:
        return ''
    # If like '4th Year' or '3rd', extract first digit
    m = re.search(r'\b([1-4])\b', s)
    if m:
        return m.group(1)
    digits = re.findall(r'\d+', s)
    if digits:
        return digits[0]
    return s[:20]

def normalize_role(val):
    """Normalize role string to standard PlayerRole choices."""
    if not val:
        return PlayerRole.BATSMAN
    s = str(val).strip().upper().replace('-', '_').replace(' ', '_')
    if 'BOWL' in s or 'BALL' in s:
        return PlayerRole.BOWLER
    if 'ALL' in s or 'ROUND' in s:
        return PlayerRole.ALL_ROUNDER
    if 'WICKET' in s or 'KEEP' in s or s == 'WK':
        return PlayerRole.WICKETKEEPER
    if 'BAT' in s:
        return PlayerRole.BATSMAN
    for choice in PlayerRole.CHOICES:
        if s == choice:
            return choice
    return PlayerRole.BATSMAN

def map_headers(raw_headers):
    """Map raw Excel/CSV headers to canonical field names."""
    mapping = {}
    for raw in raw_headers:
        if not raw:
            continue
        cleaned = re.sub(r'[_\s]+', ' ', str(raw).strip().lower())
        matched = False
        for canonical, aliases in HEADER_ALIASES.items():
            if cleaned in aliases or str(raw).strip().lower() == canonical:
                mapping[raw] = canonical
                matched = True
                break
        if not matched:
            mapping[raw] = cleaned
    return mapping

def read_rows_from_stream(content_bytes, filename=''):
    """
    Reads rows from either an .xlsx file or a .csv file.
    Returns (headers, list of raw_dicts).
    """
    is_xlsx = False
    if filename and filename.lower().endswith(('.xlsx', '.xlsm')):
        is_xlsx = True
    elif content_bytes and content_bytes[:4] == b'PK\x03\x04':
        is_xlsx = True

    if is_xlsx and HAS_OPENPYXL:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            sheet = wb.active
            rows_iter = sheet.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if not header_row:
                return [], []
            headers = [str(c).strip() if c is not None else '' for c in header_row]
            data_rows = []
            for row in rows_iter:
                if not any(c is not None and str(c).strip() for c in row):
                    continue  # skip completely blank rows
                row_dict = {}
                for idx, h in enumerate(headers):
                    if h:
                        val = row[idx] if idx < len(row) else ''
                        row_dict[h] = str(val).strip() if val is not None else ''
                data_rows.append(row_dict)
            return headers, data_rows
        except Exception as e:
            # Fallback to CSV reader if openpyxl fails
            pass

    # Read as CSV
    if isinstance(content_bytes, bytes):
        text = content_bytes.decode('utf-8-sig', errors='replace')
    else:
        text = str(content_bytes)

    stream = io.StringIO(text)
    reader = csv.DictReader(stream)
    headers = list(reader.fieldnames or [])
    data_rows = [row for row in reader]
    return headers, data_rows

def preview_players_csv(content, filename=''):
    """
    Parse and validate Excel (.xlsx) or CSV content without saving to database.
    Returns preview analysis, detected duplicates, and row-by-row status.
    """
    result = {
        'total': 0,
        'valid_count': 0,
        'duplicate_count': 0,
        'error_count': 0,
        'preview_rows': [],
        'errors': [],
        'headers_found': []
    }

    if not content:
        result['errors'].append('Uploaded file is empty.')
        return result

    content_bytes = content if isinstance(content, bytes) else content.encode('utf-8')
    headers, raw_rows = read_rows_from_stream(content_bytes, filename)

    if not headers or not raw_rows:
        result['errors'].append('File is empty or contains no readable data rows.')
        return result

    header_map = map_headers(headers)
    canonical_headers = set(header_map.values())
    result['headers_found'] = list(canonical_headers)

    if 'rule_number' not in canonical_headers or 'name' not in canonical_headers:
        result['errors'].append(
            "Missing essential columns: 'Roll Number' and 'Name' must be present in the sheet."
        )
        return result

    existing_rule_numbers = {
        (p.roll_number or '').strip().lower() for p in Player.query.all()
    }
    batch_rule_numbers = set()

    row_index = 0
    for raw_row in raw_rows:
        row_index += 1
        canonical_row = {}
        for raw_k, val in raw_row.items():
            if raw_k in header_map:
                canonical_row[header_map[raw_k]] = (str(val).strip() if val is not None else '')

        rule_no = canonical_row.get('rule_number', '').strip().upper()
        name = canonical_row.get('name', '').strip()
        photo_raw = canonical_row.get('photo', '').strip()
        year = normalize_year(canonical_row.get('year', ''))
        role = normalize_role(canonical_row.get('role', 'BATSMAN'))
        branch = canonical_row.get('branch', '').strip()
        cat = PlayerCategory.normalize(canonical_row.get('category'))

        # Photo resolution: direct URL or Google Drive link or local filename or fallback
        if is_valid_image_url(photo_raw):
            photo = normalize_photo_url(photo_raw)
        elif photo_raw:
            photo = normalize_photo_url(photo_raw)
        else:
            photo = 'default_player.png'

        row_status = 'VALID'
        status_msg = 'Ready to import'

        if not rule_no or not name:
            row_status = 'ERROR'
            status_msg = 'Missing Roll Number or Name'
            result['error_count'] += 1
        elif len(rule_no) != 10 or not re.match(r'^[A-Z0-9]{10}$', rule_no):
            row_status = 'ERROR'
            status_msg = f"Roll number '{rule_no}' must be exactly 10 alphanumeric characters"
            result['error_count'] += 1
        elif rule_no.lower() in batch_rule_numbers:
            row_status = 'DUPLICATE'
            status_msg = f"Duplicate roll number '{rule_no}' in this sheet"
            result['duplicate_count'] += 1
        elif rule_no.lower() in existing_rule_numbers:
            row_status = 'DUPLICATE'
            status_msg = f"Roll number '{rule_no}' already exists in database"
            result['duplicate_count'] += 1
        else:
            result['valid_count'] += 1
            batch_rule_numbers.add(rule_no.lower())

        result['preview_rows'].append({
            'row_num': row_index,
            'rule_number': rule_no,
            'name': name,
            'photo': photo,
            'year': year,
            'role': role,
            'branch': branch,
            'category': cat,
            'status': row_status,
            'message': status_msg
        })

    result['total'] = row_index
    return result

def parse_and_import_players_csv(content, filename=''):
    """
    Parses and imports players from Excel (.xlsx) or CSV.
    Skips invalid rows and duplicates, inserts valid players into store.
    """
    preview = preview_players_csv(content, filename)
    result = {
        'imported': 0,
        'skipped': 0,
        'duplicates': preview['duplicate_count'],
        'errors': []
    }

    if preview['errors']:
        result['errors'] = [{'row': 0, 'message': err} for err in preview['errors']]
        return result

    players_to_add = []
    for r in preview['preview_rows']:
        if r['status'] == 'VALID':
            player = Player(
                roll_number=r['rule_number'],
                name=r['name'],
                photo=r['photo'],
                year=r['year'] or None,
                role=r['role'] if r['role'] in PlayerRole.CHOICES else PlayerRole.BATSMAN,
                branch=r['branch'] or None,
                category=r['category'],
                status=PlayerStatus.AVAILABLE,
                auction_type='PRIMARY'
            )
            players_to_add.append(player)
        else:
            result['skipped'] += 1
            result['errors'].append({
                'row': r['row_num'],
                'message': f"{r['rule_number']} - {r['name']}: {r['message']}"
            })

    if players_to_add:
        try:
            db.session.add_all(players_to_add)
            db.session.commit()
            result['imported'] = len(players_to_add)
        except Exception as e:
            db.session.rollback()
            result['errors'].append({'row': 0, 'message': f"Storage error: {str(e)}"})

    return result

def export_players_excel():
    """
    Export all players to an Excel (.xlsx) file with the exact 7 columns:
    1. Roll Number
    2. Name
    3. Photo
    4. Year
    5. Role
    6. Branch
    7. Category
    """
    players = Player.query.order_by(Player.id.asc()).all()

    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SPL Players"

        # Styling definitions
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(TEMPLATE_COLUMNS)

        for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for row_idx, p in enumerate(players, start=2):
            row_data = [
                p.roll_number or '',
                p.name or '',
                p.photo or '',
                p.year or '',
                p.role or 'BATSMAN',
                p.branch or '',
                p.category or 'Rookie'
            ]
            ws.append(row_data)

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                if row_idx % 2 == 0:
                    cell.fill = alt_fill
                if col_idx in (1, 4, 5, 7):
                    cell.alignment = Alignment(horizontal='center', vertical='center')

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'SPL_Players_Roster.xlsx'

    # Fallback to CSV if openpyxl is not available
    csv_text = export_players_csv()
    return csv_text.encode('utf-8-sig'), 'text/csv', 'SPL_Players_Roster.csv'

def export_players_csv():
    """Export all players as CSV text with the exact 7 columns."""
    players = Player.query.order_by(Player.id.asc()).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TEMPLATE_COLUMNS)

    for p in players:
        writer.writerow([
            p.roll_number or '',
            p.name or '',
            p.photo or '',
            p.year or '',
            p.role or 'BATSMAN',
            p.branch or '',
            p.category or 'Rookie'
        ])

    return output.getvalue()

def generate_player_template_excel():
    """
    Generate an Excel (.xlsx) template with exact 7 columns and 3 sample rows.
    """
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Player Import Template"

        header_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        ws.append(TEMPLATE_COLUMNS)
        for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        sample_rows = [
            ['SPL26A001A', 'Arjun Reddy', 'https://images.unsplash.com/photo-1534528741775-53994a69daeb?w=300', '4', 'BATSMAN', 'CSE', 'Elite'],
            ['SPL26A002B', 'Vikram Varma', 'default_player.png', '3', 'BOWLER', 'ECE', 'Skilled'],
            ['SPL26A003C', 'Karthik Rao', 'https://images.unsplash.com/photo-1507003211169-0a1dd7228f2d?w=300', '2', 'ALL_ROUNDER', 'CSM', 'Rookie']
        ]

        for row_idx, r in enumerate(sample_rows, start=2):
            ws.append(r)
            for col_idx in range(1, len(r) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = border_thin
                if col_idx in (1, 4, 5, 7):
                    cell.alignment = Alignment(horizontal='center', vertical='center')

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 15)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'SPL_Player_Import_Template.xlsx'

    # Fallback to CSV
    csv_text = generate_player_template_csv()
    return csv_text.encode('utf-8-sig'), 'text/csv', 'SPL_Player_Import_Template.csv'

def generate_player_template_csv():
    """Generate CSV template string with 7 columns and sample rows."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerow(['SPL26A001A', 'Arjun Reddy', 'https://images.unsplash.com/photo-1534528741775-53994a69daeb?w=300', '4', 'BATSMAN', 'CSE', 'Elite'])
    writer.writerow(['SPL26A002B', 'Vikram Varma', 'default_player.png', '3', 'BOWLER', 'ECE', 'Skilled'])
    writer.writerow(['SPL26A003C', 'Karthik Rao', 'https://images.unsplash.com/photo-1507003211169-0a1dd7228f2d?w=300', '2', 'ALL_ROUNDER', 'CSM', 'Rookie'])
    return output.getvalue()
